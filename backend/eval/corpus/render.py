import re
from pathlib import Path
from typing import Literal

from docx import Document
from playwright.async_api import Page
from pydantic import BaseModel, Field, TypeAdapter

BlockKind = Literal["heading", "paragraph", "list_item", "table"]

EXCLUDED_SECTIONS = [
    "referencias",
    "notas",
    "notas y referencias",
    "enlaces externos",
    "véase también",
    "bibliografía",
    "fuentes",
]

REMOVED_SELECTORS = ", ".join(
    [
        "script",
        "style",
        "link",
        "meta",
        "img",
        "figure",
        "audio",
        "video",
        ".mw-editsection",
        ".navbox",
        ".navbox-styles",
        ".metadata",
        ".ambox",
        ".noprint",
        ".mw-empty-elt",
        ".gallery",
        ".mw-ref",
        ".mw-references-wrap",
        "ol.references",
        'a[href*="wikidata.org"]',
    ]
)

CLEANUP_SCRIPT = """
({ title, excludedSections, removedSelectors }) => {
  document.querySelectorAll(removedSelectors).forEach((node) => node.remove());
  document.querySelectorAll('h2').forEach((heading) => {
    const label = heading.textContent.trim().toLowerCase();
    if (!excludedSections.includes(label)) return;
    const section = heading.closest('section');
    if (section && section.querySelector('h1, h2, h3, h4, h5, h6') === heading) {
      section.remove();
      return;
    }
    const start = heading.parentElement.classList.contains('mw-heading')
      ? heading.parentElement
      : heading;
    let node = start.nextElementSibling;
    while (node && !node.matches('h2, .mw-heading2')) {
      const next = node.nextElementSibling;
      node.remove();
      node = next;
    }
    start.remove();
  });
  const titleHeading = document.createElement('h1');
  titleHeading.textContent = title;
  document.body.prepend(titleHeading);
}
"""

EXTRACT_BLOCKS_SCRIPT = """
() => {
  const blocks = [];
  const clean = (text) => text.replace(/\\s+/g, ' ').trim();
  const ownText = (element) => {
    const clone = element.cloneNode(true);
    clone.querySelectorAll('ul, ol').forEach((node) => node.remove());
    return clean(clone.textContent);
  };
  const visitList = (list, level) => {
    for (const item of list.children) {
      if (item.tagName.toLowerCase() !== 'li') continue;
      const text = ownText(item);
      if (text) blocks.push({ kind: 'list_item', text, level });
      item.querySelectorAll(':scope > ul, :scope > ol').forEach((nested) => visitList(nested, level + 1));
    }
  };
  const visit = (element) => {
    for (const child of element.children) {
      const tag = child.tagName.toLowerCase();
      if (/^h[1-6]$/.test(tag)) {
        const text = clean(child.textContent);
        if (text) blocks.push({ kind: 'heading', text, level: Number(tag[1]) });
      } else if (['p', 'dt', 'dd', 'pre'].includes(tag)) {
        const text = clean(child.textContent);
        if (text) blocks.push({ kind: 'paragraph', text });
      } else if (tag === 'ul' || tag === 'ol') {
        visitList(child, 1);
      } else if (tag === 'table') {
        const rows = Array.from(child.rows)
          .map((row) => Array.from(row.cells).map((cell) => clean(cell.textContent)))
          .filter((cells) => cells.some((cell) => cell));
        if (rows.length) blocks.push({ kind: 'table', rows });
      } else {
        visit(child);
      }
    }
  };
  visit(document.body);
  return blocks;
}
"""

PRINT_STYLES = """
body { font-family: 'Liberation Sans', Arial, sans-serif; font-size: 11px; line-height: 1.45; }
h1 { font-size: 22px; } h2 { font-size: 17px; } h3 { font-size: 14px; }
table { border-collapse: collapse; margin: 8px 0; }
th, td { border: 1px solid #9ca3af; padding: 2px 5px; vertical-align: top; }
a { color: inherit; text-decoration: none; }
"""


class DocumentBlock(BaseModel):
    kind: BlockKind
    text: str = ""
    level: int = 0
    rows: list[list[str]] = Field(default_factory=list)


class PreparedArticle(BaseModel):
    text: str
    blocks: list[DocumentBlock]


BLOCKS_ADAPTER = TypeAdapter(list[DocumentBlock])


def normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\xa0", " ").splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


async def prepare_article(
    page: Page, html: str, title: str, include_tables: bool = True
) -> PreparedArticle:
    await page.set_content(html, wait_until="domcontentloaded")
    await page.evaluate(
        CLEANUP_SCRIPT,
        {
            "title": title,
            "excludedSections": EXCLUDED_SECTIONS,
            "removedSelectors": REMOVED_SELECTORS if include_tables else f"{REMOVED_SELECTORS}, table",
        },
    )
    await page.add_style_tag(content=PRINT_STYLES)
    text = await page.evaluate("() => document.body.innerText")
    blocks = BLOCKS_ADAPTER.validate_python(await page.evaluate(EXTRACT_BLOCKS_SCRIPT))
    return PreparedArticle(text=normalize_text(text), blocks=blocks)


async def write_pdf(page: Page, path: Path) -> None:
    await page.pdf(
        path=str(path),
        format="A4",
        margin={"top": "18mm", "right": "14mm", "bottom": "18mm", "left": "14mm"},
    )


def write_txt(article: PreparedArticle, path: Path) -> None:
    path.write_text(article.text, encoding="utf-8")


def write_docx(article: PreparedArticle, path: Path) -> None:
    document = Document()
    for block in article.blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level - 1, 4))
        elif block.kind == "paragraph":
            document.add_paragraph(block.text)
        elif block.kind == "list_item":
            style = "List Bullet" if block.level <= 1 else f"List Bullet {min(block.level, 3)}"
            document.add_paragraph(block.text, style=style)
        else:
            column_count = max(len(row) for row in block.rows)
            table = document.add_table(rows=len(block.rows), cols=column_count)
            table.style = "Table Grid"
            for row_index, cells in enumerate(block.rows):
                for column_index, cell_text in enumerate(cells):
                    table.cell(row_index, column_index).text = cell_text
    document.save(str(path))
