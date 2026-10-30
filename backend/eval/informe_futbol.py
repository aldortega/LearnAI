# /// script
# requires-python = ">=3.13"
# dependencies = [
#     "python-docx",
# ]
# ///
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RESULTS_DIR = Path(__file__).resolve().parent / "results" / "futbol_20261018T040201Z"
FIGURES_DIR = RESULTS_DIR / "figuras"
OUTPUT_PATH = RESULTS_DIR / "informe_evaluacion_futbol.docx"

doc = Document()
figure_counter = 0
table_counter = 0


def setup_styles() -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)


def add_runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def heading(text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(paragraph, text)


def bullets(items: list[str], style: str = "List Bullet") -> None:
    for item in items:
        add_runs(doc.add_paragraph(style=style), item)


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def caption(prefix: str, number: int, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(f"{prefix} {number}. ").bold = True
    paragraph.add_run(text)


def table(headers: list[str], rows: list[list[str]], title: str, widths: list[float] | None = None) -> None:
    global table_counter
    table_counter += 1
    caption("Tabla", table_counter, title)
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(grid.rows[0].cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "31688E")
    for index, row in enumerate(rows):
        cells = grid.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = ""
            add_runs(cell.paragraphs[0], value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
            if index % 2:
                shade(cell, "EEF3F7")
    if widths:
        for row in grid.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph()


def figure(file_name: str, text: str, width_cm: float = 16) -> None:
    global figure_counter
    figure_counter += 1
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(FIGURES_DIR / file_name), width=Cm(width_cm))
    caption("Figura", figure_counter, text)


def title_page() -> None:
    title = doc.add_paragraph(style="Title")
    title.add_run("Evaluación del sistema RAG de LearnAI")
    subtitle = doc.add_paragraph()
    add_runs(subtitle, "**Informe de la evaluación con métricas RAGAS y estudio de ablación — cuaderno de fútbol**")
    doc.add_paragraph("Corrida del pipeline: futbol_20261018T040201Z · Fecha: 18 de octubre de 2026")
    doc.add_paragraph(
        "Modelo de chat: gemini-3.5-flash-lite · Embeddings: gemini-embedding-001 · "
        "Modelo juez: gemini-3.8-flash · RAGAS 0.4.3"
    )


def section_summary() -> None:
    heading("1. Resumen", 1)
    para(
        "Se evaluó el sistema de preguntas y respuestas de LearnAI, basado en recuperación aumentada "
        "(RAG), sobre un cuaderno de cinco artículos recientes de Wikipedia sobre fútbol argentino y un "
        "conjunto de 61 preguntas con respuesta de referencia. Se compararon cuatro configuraciones del "
        "pipeline en un estudio de ablación —sin recuperación, solo búsqueda vectorial, búsqueda híbrida y "
        "búsqueda híbrida con evaluación de relevancia por LLM (configuración de producción)— usando cinco "
        "métricas de RAGAS calculadas con un modelo juez y cuatro métricas deterministas de recuperación, "
        "comportamiento y latencia."
    )
    para("Los resultados principales son:")
    bullets([
        "**La recuperación es la que aporta la mejora decisiva.** La exactitud de la respuesta pasa de 0,16 "
        "sin RAG a 0,82 en producción. El RAG mejora 50 de las 55 preguntas y no empeora ninguna.",
        "**La configuración de producción recupera muy bien:** el fragmento correcto aparece entre los "
        "recuperados en el 96 % de las preguntas (Hit@k), con precisión del contexto 0,97 y exhaustividad del "
        "contexto 0,97.",
        "**Las respuestas son fieles a las fuentes:** fidelidad de 0,94. En 51 de 53 respuestas evaluadas, la "
        "respuesta es a la vez fiel y correcta.",
        "**El componente léxico de la búsqueda híbrida no muestra una mejora medible** frente a la búsqueda "
        "solo vectorial en este corpus.",
        "**La evaluación de relevancia por LLM (grading) mejora la recuperación pero no la exactitud final.** "
        "Sube la precisión del contexto de 0,79 a 0,97, ubica antes el fragmento correcto (MRR de 0,69 a 0,80), "
        "entrega al modelo 3,5 veces menos fragmentos y detecta el 100 % de las preguntas fuera del alcance. "
        "A cambio, duplica la latencia (de 2,5 s a 5,1 s de mediana) y, en dos preguntas, descarta todos los "
        "fragmentos y provoca una respuesta incorrecta.",
        "**Las preguntas de varios saltos son el punto más débil:** baja la exhaustividad del contexto (0,89) y "
        "la exactitud (0,78).",
    ])
    para(
        "El costo total de la evaluación fue de 8,11 USD en la API de Gemini (1,30 USD de generación de "
        "respuestas y 6,81 USD del modelo juez)."
    )


def section_objectives() -> None:
    heading("2. Objetivos de la evaluación", 1)
    para(
        "La evaluación busca responder, con evidencia cuantitativa, si el sistema RAG de LearnAI responde "
        "preguntas sobre las fuentes del usuario de manera correcta y fundamentada, y qué aporta cada componente "
        "del pipeline de recuperación. Concretamente, se plantearon las siguientes preguntas de evaluación:"
    )
    bullets([
        "**P1 — Aporte del RAG.** ¿Cuánto mejoran las respuestas cuando el modelo recibe fragmentos de las "
        "fuentes frente a responder solo con su conocimiento previo?",
        "**P2 — Calidad de la recuperación.** ¿El sistema encuentra los fragmentos que contienen la respuesta y "
        "los ubica en las primeras posiciones?",
        "**P3 — Fidelidad.** ¿Las respuestas se apoyan en el contexto recuperado o el modelo agrega información "
        "no sustentada?",
        "**P4 — Aporte de cada componente.** ¿Qué agregan la búsqueda léxica y la evaluación de relevancia con "
        "LLM respecto de una búsqueda vectorial simple?",
        "**P5 — Preguntas fuera del alcance.** ¿El sistema reconoce cuándo la respuesta no está en las fuentes y "
        "lo avisa al usuario?",
        "**P6 — Costo operativo.** ¿Qué latencia agrega cada componente?",
    ], style="List Number")


def section_system() -> None:
    heading("3. Sistema evaluado", 1)
    para(
        "LearnAI permite al usuario subir documentos propios (PDF, DOCX o TXT) a un cuaderno y hacer preguntas "
        "sobre ellos. El pipeline evaluado es el mismo que usa el chat de la aplicación, invocado directamente "
        "desde el script de evaluación sin pasar por la interfaz web."
    )
    heading("3.1 Indexación", 2)
    para(
        "Cada documento se convierte a texto y se divide en fragmentos respetando su estructura: los párrafos se "
        "agrupan bajo el último título detectado, las secciones con menos de 80 palabras se fusionan con la "
        "anterior y las de más de 320 se dividen en ventanas de 220 palabras con 40 de solapamiento, anteponiendo "
        "el título de la sección a cada ventana. Cada fragmento se convierte en un vector con gemini-embedding-001 "
        "y se guarda en Qdrant junto con el identificador del documento y del fragmento."
    )
    heading("3.2 Consulta", 2)
    bullets([
        "**Búsqueda vectorial con sobre-recuperación:** se piden a Qdrant 32 candidatos (4 × top_k, con top_k = 8).",
        "**Re-ranking híbrido:** cada candidato recibe un puntaje combinado 0,8 × similitud semántica + 0,2 × "
        "coincidencia léxica (fracción de palabras de la pregunta presentes en el fragmento).",
        "**Selección por umbrales y diversidad:** se admiten fragmentos con similitud ≥ 0,65 (o ≥ 0,45 con "
        "evidencia léxica) y como máximo dos por documento, hasta completar ocho.",
        "**Evaluación de relevancia con LLM (grading):** en lugar de los umbrales, el modelo de chat recibe hasta "
        "20 candidatos y marca cuáles ayudan directamente a responder. Solo se conservan los marcados como "
        "relevantes. Si no queda ninguno, el contexto queda vacío.",
        "**Generación:** el modelo gemini-3.5-flash-lite recibe los fragmentos numerados y la pregunta, y responde "
        "citando el número de cada fragmento usado.",
        "**Respaldo sin contexto:** si el contexto está vacío o es insuficiente, el modelo responde con "
        "conocimiento general y la respuesta se marca con un aviso visible para el usuario.",
    ])
    heading("3.3 Modelos utilizados", 2)
    para(
        "La evaluación involucra cuatro modelos con roles distintos. Se separó deliberadamente el modelo que "
        "**genera** las respuestas del modelo que las **evalúa**, y se eligió como juez un modelo más capaz que el "
        "generador, una práctica habitual para que el evaluador pueda detectar errores que el generador no detecta."
    )
    table(
        ["Modelo", "Proveedor", "Rol", "Configuración"],
        [
            ["gemini-3.5-flash-lite", "Google", "Modelo de chat de LearnAI: genera las respuestas y, en la configuración con grading, decide qué fragmentos son relevantes.", "Temperatura 0,2 para generar; 0 para el grading."],
            ["gemini-embedding-001", "Google", "Convierte fragmentos y preguntas en vectores para la búsqueda semántica. También lo usa RAGAS para la métrica de relevancia de la respuesta.", "Mismo modelo en indexación, consulta y evaluación."],
            ["gemini-3.8-flash", "Google", "Modelo juez de RAGAS: calcula fidelidad, relevancia, precisión y exhaustividad del contexto, y exactitud de la respuesta.", "Temperatura 0; hasta 8192 tokens de salida."],
            ["Claude", "Anthropic", "Asistente para redactar el borrador de las preguntas y respuestas de referencia, luego revisadas por el autor. No participa en la generación ni en la evaluación.", "—"],
        ],
        "Modelos utilizados y su rol en la evaluación.",
        [3.3, 2.0, 7.0, 3.7],
    )
    para(
        "gemini-3.5-flash-lite es el modelo configurado en producción: es rápido y de bajo costo, adecuado para un "
        "chat de estudio con muchas consultas. Usar como redactor de preguntas un modelo de otro proveedor (Claude) "
        "reduce el riesgo de que las preguntas queden redactadas a la medida del modelo evaluado."
    )


def section_corpus() -> None:
    heading("4. Metodología", 1)
    para(
        "La evaluación se organizó en etapas: construcción del corpus, creación del conjunto de preguntas, "
        "ejecución del pipeline en cada configuración, cálculo de métricas con el modelo juez y análisis de "
        "resultados. Cada etapa se implementó como un script reproducible dentro de `backend/eval/`."
    )
    heading("4.1 Corpus", 2)
    para(
        "Se eligieron artículos de Wikipedia en español sobre fútbol argentino de 2025 y 2026, por tres motivos. "
        "Primero, son **posteriores al corte de conocimiento del modelo** (enero de 2025), de modo que el modelo no "
        "puede responder correctamente de memoria y la evaluación mide de verdad el aporte de la recuperación. "
        "Segundo, su licencia (CC BY-SA 4.0) permite redistribuirlos como parte de la tesis. Tercero, se priorizaron "
        "artículos con predominio de texto sobre tablas, que representan mejor el uso esperado de la plataforma."
    )
    para(
        "El script `fetch_wikipedia.py` fija una revisión concreta de cada artículo (reproducibilidad), elimina "
        "referencias, imágenes, cajas de navegación y secciones de bibliografía, y genera los archivos en los tres "
        "formatos que admite la plataforma (PDF, DOCX y TXT) a partir del mismo contenido depurado, para ejercitar "
        "también los distintos caminos de ingesta."
    )
    table(
        ["Artículo", "Formato", "Revisión", "Caracteres", "Preguntas que lo usan"],
        [
            ["Final de la Copa Mundial de Fútbol de 2026", "DOCX", "175872794", "39 409", "22"],
            ["Claudio Tapia", "DOCX", "175657666", "9 416", "12"],
            ["El último tango", "PDF", "175801639", "6 737", "11"],
            ["Independiente vs. Universidad de Chile (Copa Sudamericana 2025)", "TXT", "173383031", "14 288", "9"],
            ["Pasillo-gate", "PDF", "173886104", "7 884", "8"],
        ],
        "Documentos del cuaderno de fútbol (77 734 caracteres en total). Una pregunta de varios saltos puede usar "
        "más de un documento.",
        [7.5, 1.8, 2.2, 2.0, 2.5],
    )


def section_dataset() -> None:
    heading("4.2 Conjunto de preguntas", 2)
    para(
        "Se construyó un conjunto de 61 preguntas en español. Para cada una se registró la pregunta, su tipo, "
        "una **respuesta de referencia**, los **pasajes del texto** que la sustentan y los **identificadores de los "
        "fragmentos indexados** que contienen esa información (exportados desde Qdrant con `export_chunks.py`). "
        "Las preguntas fueron redactadas con asistencia de un modelo de lenguaje (Claude) a partir de los "
        "fragmentos indexados y luego revisadas manualmente por el autor; solo las preguntas marcadas como "
        "revisadas entran en la evaluación."
    )
    table(
        ["Tipo", "Cantidad", "Descripción"],
        [
            ["Un solo salto (single_hop)", "45", "La respuesta está en un único pasaje del corpus."],
            ["Varios saltos (multi_hop)", "10", "La respuesta requiere combinar información de dos o más fragmentos, a veces de documentos distintos."],
            ["Fuera del alcance (out_of_scope)", "6", "La respuesta no está en el cuaderno. Se espera que el sistema lo advierta."],
        ],
        "Composición del conjunto de preguntas.",
        [4.5, 2.0, 9.5],
    )
    table(
        ["ID", "Tipo", "Pregunta", "Respuesta de referencia"],
        [
            ["fut-001", "Un salto", "¿Cuándo fue elegido Claudio Tapia presidente de la AFA y con cuántos votos?", "El 29 de marzo de 2017, con 40 votos a favor y 3 abstenciones."],
            ["fut-052", "Varios saltos", "¿Cómo le fue a Cabo Verde contra España y contra Argentina en el Mundial 2026?", "Empató 0:0 con España en la fase de grupos […]. Contra Argentina, en dieciseisavos, perdió 3:2 en tiempo extra."],
            ["fut-058", "Fuera del alcance", "¿Qué equipo ganó la Copa Sudamericana 2025?", "La información no está en las fuentes del cuaderno."],
        ],
        "Ejemplos del conjunto de preguntas.",
        [1.6, 2.2, 6.0, 6.2],
    )
    para(
        "Durante la redacción se detectaron contradicciones internas en las fuentes (por ejemplo, el minuto de un "
        "gol informado de dos formas distintas en el mismo artículo o nombres escritos de manera diferente). Esas "
        "preguntas se descartaron o se reformularon para que la referencia no fuera ambigua."
    )


def section_configs() -> None:
    heading("4.3 Estudio de ablación", 2)
    para(
        "Un estudio de ablación compara variantes del sistema en las que se agregan o quitan componentes de a uno, "
        "para atribuir a cada componente su efecto. Se definieron cuatro configuraciones que forman una cadena: "
        "cada una agrega una pieza a la anterior. El resto de los parámetros se mantuvo fijo (top_k = 8, umbral "
        "0,65, máximo dos fragmentos por documento, mismo modelo y mismo prompt)."
    )
    table(
        ["Configuración", "Búsqueda vectorial", "Componente léxico", "Grading LLM", "Qué mide frente a la anterior"],
        [
            ["Sin RAG (no_rag)", "—", "—", "—", "Línea base: el modelo responde solo con conocimiento previo."],
            ["Solo vectorial (vector_only)", "Sí", "No (peso 0)", "No", "El aporte de recuperar contexto."],
            ["Híbrida (hybrid)", "Sí", "Sí (peso 0,2)", "No", "El aporte del re-ranking léxico."],
            ["Híbrida + grading (hybrid_grading)", "Sí", "Sí (peso 0,2)", "Sí", "El aporte del filtro de relevancia con LLM. Es la configuración de producción."],
        ],
        "Configuraciones del estudio de ablación.",
        [3.6, 2.0, 2.2, 1.8, 6.4],
    )
    para("A continuación se describe qué hace cada configuración y qué se espera observar en ella.")
    heading("Sin RAG (no_rag)", 3)
    para(
        "El modelo recibe solo la pregunta, sin ningún fragmento de los documentos, y responde con lo que aprendió "
        "durante su entrenamiento. No representa un uso real de la plataforma: es la **línea base** que permite medir "
        "cuánto aporta el RAG. Como el corpus trata hechos de 2025 y 2026, posteriores al corte de conocimiento del "
        "modelo, se espera que acierte poco y que muchas respuestas digan que el evento todavía no ocurrió. En esta "
        "configuración no hay contexto, por lo que todas las respuestas llevan la marca de conocimiento general y no "
        "se pueden calcular las métricas de recuperación."
    )
    heading("Solo vectorial (vector_only)", 3)
    para(
        "Es el **RAG más simple**. La pregunta se convierte en un vector y se buscan en Qdrant los fragmentos con "
        "mayor similitud de coseno, es decir, los que tienen un **significado** parecido aunque usen otras palabras. "
        "Los candidatos se ordenan solo por esa similitud, se filtran con el umbral de 0,65 y la regla de diversidad, "
        "y los ocho primeros se entregan al modelo. Es el punto de comparación para saber si los componentes "
        "adicionales justifican su complejidad."
    )
    heading("Híbrida (hybrid)", 3)
    para(
        "Agrega a la búsqueda vectorial un **componente léxico**: a cada candidato se le suma un puntaje según cuántas "
        "palabras de la pregunta aparecen literalmente en el fragmento, con un peso de 0,2 frente a 0,8 de la "
        "similitud semántica. La idea es compensar una debilidad conocida de los embeddings, que a veces no "
        "distinguen bien términos exactos como nombres propios, siglas o cifras. Además, la evidencia léxica permite "
        "admitir fragmentos con similitud semántica algo menor (≥ 0,45). Se espera una mejora en preguntas que "
        "dependen de un término específico."
    )
    heading("Híbrida + grading (hybrid_grading)", 3)
    para(
        "Es la **configuración de producción** de LearnAI. Parte de la búsqueda híbrida, pero reemplaza el filtro por "
        "umbrales por una **evaluación hecha por el propio LLM**: el modelo de chat recibe la pregunta y hasta 20 "
        "candidatos numerados, y marca uno por uno si ayuda directamente a responder. Solo los fragmentos marcados "
        "como relevantes llegan a la generación, en el orden del re-ranking. Los umbrales miden parecido; el grading "
        "intenta medir utilidad. Se espera un contexto más chico y más preciso y una mejor detección de preguntas sin "
        "respuesta en las fuentes, a cambio de una llamada adicional al LLM por consulta."
    )


def section_procedure() -> None:
    heading("4.4 Procedimiento", 2)
    bullets([
        "**Generación de respuestas** (`run_pipeline.py`): cada una de las 61 preguntas se respondió con las cuatro "
        "configuraciones y tres repeticiones, en total 732 respuestas, sin errores. Para cada respuesta se guardaron "
        "los fragmentos recuperados, las fuentes citadas, la marca de conocimiento general y la latencia.",
        "**Evaluación con el modelo juez** (`evaluate_ragas.py`): se calcularon las métricas RAGAS sobre la "
        "primera repetición de cada configuración (987 puntajes, todos válidos). Las repeticiones 2 y 3 del juez "
        "no se ejecutaron por costo; las métricas deterministas (marca de conocimiento general y latencia) sí se "
        "analizaron sobre las tres repeticiones cuando se indica.",
        "**Análisis y figuras** (`plots.py`): las figuras de este informe se generan automáticamente a partir de "
        "los resultados guardados, sin nuevas llamadas a la API.",
    ])


def section_ragas() -> None:
    heading("4.5 RAGAS y el enfoque LLM como juez", 2)
    para(
        "Evaluar un sistema RAG es difícil con métricas clásicas de texto como BLEU o ROUGE, que cuentan palabras "
        "coincidentes con una referencia: una respuesta correcta redactada de otra forma obtiene un puntaje bajo, y "
        "una respuesta que copia palabras pero cambia un dato obtiene uno alto. Tampoco miden si la respuesta se "
        "apoya en los documentos recuperados. La alternativa más precisa, la evaluación humana, es lenta y costosa, "
        "y no escala a cientos de respuestas por cada cambio del sistema."
    )
    heading("LLM como juez", 3)
    para(
        "El enfoque **LLM como juez** (LLM-as-a-judge) usa un modelo de lenguaje para evaluar las salidas de otro "
        "sistema, siguiendo instrucciones y criterios definidos. Se apoya en que los modelos actuales comprenden el "
        "significado del texto y pueden juzgar si dos respuestas dicen lo mismo o si una afirmación se deduce de un "
        "pasaje, con un acuerdo con evaluadores humanos que la literatura reporta como alto. Sus ventajas son la "
        "escala, el costo y la repetibilidad. Sus riesgos también están documentados: preferencia por respuestas más "
        "largas, preferencia por textos generados por modelos de la misma familia, sensibilidad a la redacción del "
        "prompt y cierta variabilidad entre ejecuciones."
    )
    para("Para reducir esos riesgos se aplicaron las prácticas habituales:")
    bullets([
        "**Temperatura 0** en el juez, para que sus decisiones sean lo más estables posible.",
        "**Un juez más capaz que el generador** (gemini-3.8-flash frente a gemini-3.5-flash-lite).",
        "**Tareas pequeñas y verificables** en lugar de pedir una nota global: RAGAS descompone cada evaluación en "
        "decisiones simples, como «¿esta afirmación se deduce de este texto?».",
        "**Salidas estructuradas:** el juez responde en un formato JSON con un esquema fijo, que el programa valida, "
        "en lugar de texto libre.",
        "**Respuestas de referencia escritas y revisadas por una persona**, contra las cuales se compara.",
    ])
    heading("Cómo funciona RAGAS", 3)
    para(
        "**RAGAS** (Retrieval Augmented Generation Assessment) es un framework de código abierto, propuesto por Es "
        "et al. (2023), que implementa este enfoque con métricas específicas para sistemas RAG. Su idea central es "
        "evaluar por separado las dos partes del sistema: el **recuperador** (¿trajo los fragmentos correctos?) y el "
        "**generador** (¿usó bien esos fragmentos para responder?). Así, cuando una respuesta falla, se puede saber "
        "en qué etapa se originó el problema."
    )
    para(
        "Cada muestra que se evalúa contiene hasta cuatro elementos: la **pregunta**, la **respuesta** del sistema, los "
        "**fragmentos recuperados** y la **respuesta de referencia**. Cada métrica usa una combinación distinta de esos "
        "elementos y sigue el mismo esquema general:"
    )
    bullets([
        "**Descomponer:** el juez divide un texto (la respuesta o la referencia) en afirmaciones simples e "
        "independientes.",
        "**Verificar:** para cada afirmación o fragmento, el juez toma una decisión binaria (sustentado o no, útil o "
        "no) y explica brevemente por qué.",
        "**Agregar:** el programa, no el juez, convierte esas decisiones en un puntaje entre 0 y 1 con una fórmula "
        "fija (por ejemplo, la proporción de afirmaciones sustentadas).",
    ], style="List Number")
    para(
        "Por ejemplo, para calcular la fidelidad de la respuesta «Tapia fue elegido en 2017 con 40 votos», el juez "
        "la separa en «Tapia fue elegido presidente de la AFA en 2017» y «Tapia obtuvo 40 votos», verifica cada una "
        "contra los fragmentos recuperados, y si ambas están sustentadas la fidelidad es 2/2 = 1. Este esquema hace "
        "que el puntaje sea interpretable: se puede revisar qué afirmación concreta falló."
    )
    para(
        "Algunas métricas combinan el juez con otras técnicas. La relevancia de la respuesta, por ejemplo, pide al "
        "juez que genere las preguntas que la respuesta estaría contestando y luego mide con embeddings cuánto se "
        "parecen a la pregunta original, sin comparar con una referencia."
    )
    table(
        ["Métrica", "Evalúa", "Usa referencia", "Usa contexto"],
        [
            ["Fidelidad", "Generador", "No", "Sí"],
            ["Relevancia de la respuesta", "Generador", "No", "No"],
            ["Exactitud de la respuesta", "Sistema completo", "Sí", "No"],
            ["Precisión del contexto", "Recuperador", "Sí", "Sí"],
            ["Exhaustividad del contexto", "Recuperador", "Sí", "Sí"],
        ],
        "Qué parte del sistema evalúa cada métrica de RAGAS y qué datos necesita.",
        [5.0, 4.0, 3.0, 3.0],
    )
    para(
        "En este trabajo se usó RAGAS 0.4.3 con su API de métricas más reciente (`ragas.metrics.collections`). El "
        "juez se conectó a través del endpoint de Gemini compatible con la API de OpenAI, y el script de evaluación "
        "envía las muestras en paralelo (cuatro a la vez), guarda cada puntaje apenas se calcula y reintenta los que "
        "fallan, de modo que una interrupción no obliga a repetir lo ya evaluado."
    )


def section_metrics() -> None:
    heading("4.6 Métricas", 2)
    para(
        "Se usaron dos familias de métricas. Las **métricas RAGAS** se calculan con un modelo de lenguaje que actúa "
        "como juez y toman valores entre 0 y 1 (mayor es mejor). Las **métricas deterministas** se calculan "
        "directamente a partir de los resultados, sin juez."
    )
    table(
        ["Métrica", "Qué responde", "Cómo se calcula", "Requiere"],
        [
            ["Fidelidad (Faithfulness)", "¿La respuesta se apoya en el contexto recuperado?", "El juez descompone la respuesta en afirmaciones y verifica cuáles se deducen del contexto. Puntaje = afirmaciones sustentadas / total.", "Respuesta, contexto"],
            ["Relevancia de la respuesta (Answer Relevancy)", "¿La respuesta atiende a lo que se preguntó?", "El juez genera preguntas a partir de la respuesta; se promedia su similitud coseno (embeddings) con la pregunta original.", "Pregunta, respuesta"],
            ["Precisión del contexto (Context Precision)", "¿Los fragmentos recuperados son útiles y están bien ordenados?", "El juez decide si cada fragmento es útil para la respuesta de referencia; se calcula la precisión promedio ponderada por posición.", "Contexto, referencia"],
            ["Exhaustividad del contexto (Context Recall)", "¿El contexto contiene toda la información necesaria?", "Fracción de las afirmaciones de la referencia que pueden atribuirse al contexto recuperado.", "Contexto, referencia"],
            ["Exactitud de la respuesta (Answer Accuracy)", "¿La respuesta coincide con la referencia?", "Dos evaluaciones independientes del juez califican la concordancia (0, 0,5 o 1) y se promedian. Valores posibles: 0; 0,25; 0,5; 0,75; 1.", "Respuesta, referencia"],
            ["Hit@k", "¿Se recuperó al menos un fragmento de referencia?", "1 si algún fragmento de referencia está entre los recuperados; 0 si no.", "Fuentes, fragmentos de referencia"],
            ["MRR", "¿En qué posición aparece el primer fragmento correcto?", "1 / posición del primer fragmento de referencia (0 si no aparece).", "Fuentes, fragmentos de referencia"],
            ["Conocimiento general", "¿El sistema avisó que respondió fuera de las fuentes?", "Porcentaje de respuestas que llevan la marca de conocimiento general.", "Respuesta"],
            ["Latencia", "¿Cuánto tarda una consulta completa?", "Tiempo desde la recuperación hasta la respuesta generada.", "—"],
        ],
        "Métricas utilizadas.",
        [3.3, 3.6, 6.4, 2.7],
    )
    para(
        "Las métricas RAGAS se calculan solo sobre las 55 preguntas dentro del alcance. Las preguntas fuera del "
        "alcance se analizan con la marca de conocimiento general. En la configuración sin RAG no hay contexto, "
        "por lo que solo aplican la relevancia y la exactitud de la respuesta."
    )
    heading("4.7 Decisiones de diseño de la evaluación", 2)
    bullets([
        "**Juez:** gemini-3.8-flash, temperatura 0, hasta 8192 tokens de salida (el modelo usa razonamiento interno "
        "y el límite por defecto truncaba respuestas). Se accedió mediante el endpoint compatible con OpenAI de "
        "Gemini, usando RAGAS 0.4.3. El script corre en un entorno aislado (PEP 723) para no alterar las "
        "dependencias de la aplicación.",
        "**Exactitud de la respuesta en lugar de Factual Correctness.** En pruebas iniciales, Factual Correctness "
        "asignaba 0 a respuestas correctas pero más detalladas que la referencia, porque penaliza cada afirmación "
        "extra como falso positivo. Por ejemplo, fut-006 obtuvo 0 con Factual Correctness y 1 con Answer Accuracy, "
        "y fut-041 obtuvo 0 frente a 0,75. Answer Accuracy compara el contenido esencial y es más adecuada para "
        "respuestas de un asistente que tienden a dar contexto adicional.",
        "**Reintentos:** los puntajes que fallaron por tiempo de espera (cuatro, en una pregunta) se recalcularon; "
        "el resultado final no tiene puntajes faltantes.",
    ])


def section_results_overview() -> None:
    heading("5. Resultados", 1)
    heading("5.1 Visión general", 2)
    para(
        f"La Tabla {table_counter + 1} y la Figura {figure_counter + 1} resumen todas las métricas para las "
        f"preguntas dentro del alcance. La Figura {figure_counter + 2} muestra las métricas RAGAS como gráfico de "
        "barras."
    )
    table(
        ["Configuración", "Fidelidad", "Relevancia", "Precisión ctx.", "Exhaustividad ctx.", "Exactitud", "Hit@k", "MRR"],
        [
            ["Sin RAG", "—", "0,76", "—", "—", "0,16", "—", "—"],
            ["Solo vectorial", "**0,97**", "**0,90**", "0,79", "0,92", "**0,83**", "0,91", "0,69"],
            ["Híbrida", "0,94", "0,90", "0,79", "0,92", "0,80", "0,89", "0,69"],
            ["Híbrida + grading", "0,94", "0,87", "**0,97**", "**0,97**", "0,82", "**0,96**", "**0,80**"],
        ],
        "Resultados por configuración (preguntas dentro del alcance, n = 55; repetición 1).",
        [3.4, 1.7, 1.7, 1.8, 2.1, 1.7, 1.5, 1.5],
    )
    figure("01b_mapa_de_calor.png", "Métricas por configuración. Las celdas vacías corresponden a métricas que requieren contexto recuperado.")
    figure("01_metricas_ragas.png", "Métricas RAGAS por configuración (preguntas dentro del alcance, n = 55).")
    figure("11_radar_configuraciones.png", "Métricas RAGAS medias por configuración (n = 55). Sin RAG solo se muestra en las métricas que no requieren contexto recuperado.", 11)
    para(
        "A simple vista se distinguen dos patrones. Primero, la línea base sin RAG queda muy por debajo en "
        "exactitud. Segundo, entre las configuraciones con RAG, las métricas de la respuesta (fidelidad, relevancia, "
        "exactitud) son muy parecidas, mientras que las de recuperación (precisión y exhaustividad del contexto, "
        "Hit@k, MRR) separan claramente a la configuración con grading."
    )
    para(
        f"El gráfico de radar (Figura {figure_counter}) permite verlo de un vistazo: los polígonos de las tres configuraciones con RAG "
        "se superponen casi por completo en fidelidad, relevancia y exactitud, y solo la configuración con grading "
        "se expande hacia afuera en precisión y exhaustividad del contexto. Sin RAG no se puede trazar un polígono "
        "porque tres de las cinco métricas requieren contexto recuperado; se muestran solo los dos puntos aplicables."
    )


def section_results_rag() -> None:
    heading("5.2 Aporte del RAG (P1)", 2)
    para(
        "Sin recuperación, el modelo obtiene una exactitud media de 0,16: el 69 % de sus respuestas recibe 0 y "
        "ninguna alcanza el puntaje máximo. Con la configuración de producción, la exactitud sube a 0,82 y el 87 % "
        f"de las respuestas obtiene 0,75 o 1. La comparación pregunta por pregunta (Figura {figure_counter + 1}) muestra que el RAG "
        "mejora 50 de las 55 preguntas, empata en 5 y no empeora ninguna: ningún punto queda debajo de la diagonal."
    )
    figure("02b_exactitud_por_pregunta_puntos.png", "Exactitud por pregunta sin RAG frente a la configuración de producción. Cada punto es una pregunta; los puntos con igual puntaje se agrupan alrededor del valor.", 11)
    figure("09_distribucion_exactitud.png", "Distribución de los puntajes de exactitud de la respuesta por configuración (n = 55).")
    para(
        "Es interesante que la relevancia de la respuesta sin RAG sea relativamente alta (0,76). El modelo responde "
        "sobre el tema preguntado, pero con datos incorrectos o declarando que el evento aún no ocurrió (su "
        "conocimiento termina en enero de 2025). Esto ilustra por qué la relevancia no alcanza por sí sola como "
        "medida de calidad y hace falta una métrica de exactitud contra una referencia."
    )


def section_results_retrieval() -> None:
    heading("5.3 Calidad de la recuperación (P2)", 2)
    para(
        f"La curva Hit@k (Figura {figure_counter + 1}) indica, para cada posición k, el porcentaje de preguntas en las que algún "
        "fragmento de referencia aparece entre los k primeros recuperados. Con grading, el fragmento correcto queda "
        "en primer lugar en el 67 % de las preguntas (frente a 55-56 % sin grading) y entre los tres primeros en el "
        "95 %. Sin grading, las curvas de la búsqueda vectorial y la híbrida son prácticamente iguales y llegan al "
        "89-91 % recién con k = 8."
    )
    figure("10_curva_hit_k.png", "Porcentaje de preguntas con al menos un fragmento de referencia entre los k primeros recuperados.")
    para(
        "La precisión del contexto confirma el mismo efecto desde otro ángulo: sin grading, el modelo recibe en "
        "promedio 7,5 fragmentos, de los cuales muchos no son útiles (precisión 0,79). Con grading recibe 2,1 "
        "fragmentos, casi todos útiles (precisión 0,97). La exhaustividad también sube, de 0,92 a 0,97, "
        "probablemente porque el grading evalúa 20 candidatos en lugar de filtrar por umbral y rescata fragmentos "
        "útiles que los umbrales dejaban afuera."
    )


def section_results_ablation() -> None:
    heading("5.4 Aporte de cada componente (P4)", 2)
    para(
        f"La Figura {figure_counter + 1} compara pregunta por pregunta cada paso de la ablación: a la derecha, las preguntas que mejoran "
        "al agregar el componente; a la izquierda, las que empeoran; al margen, las que no cambian."
    )
    figure("06b_comparacion_divergente.png", "Preguntas que mejoran (derecha) o empeoran (izquierda) en cada paso de la ablación.")
    bullets([
        "**Recuperar contexto (sin RAG → solo vectorial):** 49 preguntas mejoran y ninguna empeora. Es, con "
        "diferencia, el paso de mayor impacto.",
        "**Agregar el componente léxico (solo vectorial → híbrida):** no hay una tendencia. En exactitud mejoran 4 "
        "preguntas y empeoran 7; en precisión del contexto, 11 y 13. Las diferencias se compensan y las medias son "
        "casi iguales. Una explicación plausible es que las preguntas están formuladas en lenguaje natural y "
        "parafrasean el texto, un escenario donde los embeddings ya funcionan bien. El componente léxico, con peso "
        "0,2, debería ayudar sobre todo con términos exactos (siglas, códigos, nombres poco frecuentes), que en este "
        "corpus rara vez son la única pista.",
        "**Agregar grading (híbrida → híbrida + grading):** la precisión del contexto mejora en 25 preguntas y "
        "empeora en 1. En exactitud el efecto es pequeño y mixto (8 mejoran, 5 empeoran): el modelo de chat ya "
        "sabía ignorar los fragmentos irrelevantes cuando recibía ocho.",
    ])
    heading("Diferencia entre híbrida e híbrida + grading", 3)
    para(
        "Como el grading es el único componente cuyo aporte depende de qué métrica se mire, se comparó con más "
        "detalle contra la búsqueda híbrida sin grading. Para cada métrica RAGAS se calculó la diferencia por "
        "pregunta (con grading menos sin grading) y, a partir de esas diferencias pareadas, la media y su intervalo "
        "de confianza del 95 % mediante bootstrap (10 000 remuestreos de las preguntas). Un intervalo que no incluye "
        "el cero indica una diferencia sistemática; uno que lo incluye es compatible con ausencia de efecto. Las "
        "métricas que requieren contexto se calculan sobre 53 preguntas, porque en dos de ellas el grading deja el "
        "contexto vacío (sección 6)."
    )
    table(
        ["Métrica", "n", "Híbrida", "Híbrida + grading", "Diferencia", "IC 95 %", "Mejoran / empeoran"],
        [
            ["Precisión del contexto", "53", "0,80", "0,97", "**+0,172**", "[+0,106; +0,244]", "25 / 1"],
            ["Exhaustividad del contexto", "53", "0,91", "0,97", "**+0,063**", "[+0,013; +0,126]", "6 / 1"],
            ["Exactitud de la respuesta", "55", "0,80", "0,82", "+0,018", "[−0,050; +0,086]", "8 / 5"],
            ["Fidelidad", "53", "0,94", "0,94", "+0,006", "[−0,058; +0,064]", "6 / 4"],
            ["Relevancia de la respuesta", "55", "0,90", "0,87", "−0,022", "[−0,073; +0,025]", "22 / 28"],
        ],
        "Diferencia media por pregunta entre híbrida + grading e híbrida, con intervalo de confianza bootstrap del "
        "95 %. En negrita, las diferencias cuyo intervalo no incluye el cero.",
        [3.8, 0.9, 1.7, 2.2, 1.9, 3.1, 2.4],
    )
    figure("12_diferencia_grading_ic.png", "Diferencia media por pregunta entre híbrida + grading e híbrida, con intervalo de confianza del 95 % (bootstrap). Los intervalos en color no incluyen el cero.")
    para(
        "El resultado es nítido. Las dos métricas de recuperación mejoran de forma sistemática: la precisión del "
        "contexto en 0,17 puntos y la exhaustividad en 0,06. En cambio, las tres métricas de la respuesta final "
        "(exactitud, fidelidad y relevancia) tienen intervalos que incluyen el cero: con esta muestra no hay "
        "evidencia de que el grading las mejore ni las empeore. Esto respalda estadísticamente la conclusión de la "
        "sección anterior, que antes se apoyaba solo en conteos de preguntas."
    )
    figure("13_hibrida_vs_grading_pareado.png", "Puntaje por pregunta con y sin grading. Los puntos sobre la diagonal empatan; por encima mejora con grading y por debajo empeora. El número indica la cantidad de preguntas.")
    para(
        "La Figura "
        f"{figure_counter} muestra las mismas diferencias pregunta por pregunta. En exactitud, la gran mayoría "
        "de las preguntas empata (23 en 1,0 y 15 en 0,75), 8 mejoran y 5 empeoran; entre estas últimas hay dos que "
        "caen a 0, que probablemente corresponden al modo de falla en que el grading descarta todos los fragmentos. En precisión "
        "del contexto, casi todos los puntos quedan sobre la diagonal, es decir, el grading mejora la precisión de "
        "manera generalizada y no por unos pocos casos."
    )
    para(
        "En síntesis, el grading no mejora la exactitud de las respuestas, pero sí la calidad del contexto, el orden "
        "de las fuentes citadas, la detección de preguntas fuera del alcance (sección 5.7) y la cantidad de "
        "contexto que se envía al modelo. Este último punto reduce los tokens de entrada de la generación, aunque se "
        "suma una llamada extra al LLM para el grading."
    )


def section_results_faithfulness() -> None:
    heading("5.5 Fidelidad (P3)", 2)
    para(
        f"La Figura {figure_counter + 1} cruza fidelidad y exactitud para cada pregunta de la configuración de producción. 51 de las 53 "
        "respuestas con contexto son a la vez fieles y correctas. No hay respuestas fieles pero incorrectas, lo que "
        "indica que cuando el sistema falla no es porque copie mal el contexto. Tampoco hay respuestas infieles e "
        "incorrectas, que serían el patrón típico de una alucinación."
    )
    figure("08_fidelidad_vs_exactitud.png", "Fidelidad frente a exactitud por pregunta en la configuración de producción. Se agregó un desplazamiento vertical pequeño para evitar superposiciones.", 12)
    figure("08b_fidelidad_vs_exactitud_por_config.png", "Fidelidad frente a exactitud por pregunta en cada configuración con recuperación. Se agregó un desplazamiento pequeño para evitar superposiciones.")
    para(
        "Al repetir el gráfico para las tres configuraciones con recuperación se ve que el patrón es el mismo: casi "
        "todas las preguntas se concentran en fidelidad cercana a 1, y la exactitud queda repartida en los niveles "
        "0,5, 0,75 y 1. Los puntos de baja fidelidad son pocos en las tres configuraciones."
    )
    para(
        "Las dos respuestas «correctas pero no sustentadas» (fut-024 y fut-039) son casos de información implícita: "
        "el fragmento recuperado contiene el dato, pero no menciona explícitamente a los protagonistas porque eso "
        "se dice en otro párrafo. El juez no puede verificar la afirmación completa solo con ese fragmento y "
        "asigna fidelidad 0. Con la búsqueda vectorial (ocho fragmentos), esos mismos casos obtienen fidelidad 1, "
        "porque el contexto más amplio incluye la información que faltaba. Esto explica la fidelidad levemente "
        "menor de la configuración de producción (0,94 frente a 0,97 de solo vectorial)."
    )


def section_results_hops() -> None:
    heading("5.6 Preguntas de un salto y de varios saltos", 2)
    figure("04_salto_simple_vs_multiple.png", "Configuración de producción: preguntas de un solo salto (n = 45) frente a preguntas de varios saltos (n = 10).")
    para(
        "Las preguntas de varios saltos son más difíciles para la recuperación: la exhaustividad del contexto baja de "
        "0,99 a 0,89, Hit@k de 0,98 a 0,90 y MRR de 0,82 a 0,70, porque basta con que falte uno de los fragmentos "
        "necesarios para que la respuesta quede incompleta. Eso se traslada a la exactitud, que baja de 0,83 a 0,78. "
        "La fidelidad, en cambio, es incluso mayor (0,98): cuando el sistema responde, lo hace con lo que recuperó. "
        "Hay que tener en cuenta que el grupo de varios saltos tiene solo 10 preguntas, por lo que estas diferencias "
        "son indicativas."
    )


def section_results_scope() -> None:
    heading("5.7 Preguntas fuera del alcance (P5)", 2)
    para(
        "El sistema marca una respuesta como conocimiento general cuando no encontró contexto suficiente. Lo deseable "
        "es que esa marca aparezca en todas las preguntas fuera del alcance y en pocas de las que sí tienen respuesta "
        "en las fuentes. Esta métrica se calculó sobre las tres repeticiones del pipeline."
    )
    figure("05_conocimiento_general.png", "Porcentaje de respuestas marcadas como conocimiento general. Media y desvío estándar sobre 3 corridas del pipeline.")
    bullets([
        "**Con grading se detecta el 100 % de las preguntas fuera del alcance**, frente al 83 % de la búsqueda "
        "vectorial y el 78 % (± 10 %) de la híbrida. Sin grading, los umbrales de similitud dejan pasar fragmentos "
        "del mismo tema que no responden la pregunta, y el modelo intenta responder con ellos.",
        "**El costo es un 10 % de falsas alarmas dentro del alcance**, frente al 5 % sin grading. Corresponden en "
        "parte a los casos en que el grading descarta todos los fragmentos (sección 6).",
    ])
    para(
        "Un aspecto a destacar es qué hace el modelo después de avisar. En fut-058 («¿Qué equipo ganó la Copa "
        "Sudamericana 2025?») el sistema advierte correctamente que la información no está en las fuentes, pero "
        "luego responde de memoria que el torneo «aún no se ha disputado». El aviso protege al usuario, que sabe que "
        "esa parte no proviene de sus documentos, pero el contenido puede contradecir información real."
    )


def section_results_cost() -> None:
    heading("5.8 Latencia y costo (P6)", 2)
    figure("07_latencia_boxplot.png", "Distribución de la latencia por consulta y configuración (3 corridas del pipeline, todas las preguntas).")
    para(
        "La mediana de latencia es de 2,1 s sin RAG, 2,8 s con búsqueda vectorial, 2,5 s con búsqueda híbrida y 5,1 s "
        "con grading. El grading agrega unos 2,5 s de forma estable (toda la distribución se desplaza), porque suma una "
        "llamada al LLM antes de generar la respuesta. Los valores atípicos de 10 a 14 s aparecen en todas las "
        "configuraciones y corresponden a demoras ocasionales de la API, no a un componente en particular. La "
        "diferencia entre la búsqueda vectorial y la híbrida no es relevante: el cálculo léxico es local y su costo "
        "es despreciable."
    )
    figure("03_calidad_vs_latencia.png", "Exactitud de la respuesta frente a la latencia mediana por configuración. El tamaño del círculo indica la cantidad media de fragmentos entregados al modelo.")
    para(
        "Para mirar la cola de la distribución, la Figura "
        f"{figure_counter + 1} muestra la mediana, el percentil 95 y el máximo de cada configuración."
    )
    figure("07c_latencia_percentiles.png", "Latencia por consulta: mediana, percentil 95 y máximo por configuración (3 corridas del pipeline, todas las preguntas).")
    table(
        ["Configuración", "Mediana (s)", "p95 (s)", "Máximo (s)"],
        [
            ["Sin RAG", "2,1", "3,7", "13,3"],
            ["Solo vectorial", "2,8", "4,9", "14,2"],
            ["Híbrida", "2,5", "4,5", "10,9"],
            ["Híbrida + grading", "5,1", "8,0", "10,5"],
        ],
        "Latencia por consulta (todas las preguntas, 3 corridas del pipeline).",
        [4.5, 3.0, 3.0, 3.0],
    )
    para(
        "El grading duplica la mediana y casi duplica el percentil 95 (de 4,5 s a 8,0 s). Los máximos de las demás "
        "configuraciones son picos aislados de la API y no reflejan el comportamiento típico; el aumento del grading, "
        "en cambio, se ve en toda la distribución."
    )
    heading("Costo y beneficio del grading", 3)
    para(
        f"La Figura {figure_counter + 1} combina latencia y calidad: para cada configuración muestra cuántos "
        "segundos extra de latencia mediana y cuánta exactitud adicional aporta frente a no usar RAG."
    )
    figure("03b_costo_beneficio.png", "Ganancia de exactitud de la respuesta frente a la latencia mediana extra, tomando como referencia la configuración sin RAG (n = 55).")
    para(
        "Las tres configuraciones con recuperación logran prácticamente la misma ganancia de exactitud (entre +0,64 "
        "y +0,67), pero con costos distintos: la búsqueda híbrida agrega unos 0,4 s, la vectorial 0,6 s y el "
        "grading unos 3 s (2,6 s más que la búsqueda híbrida sola). Con estos datos, el tiempo adicional del grading no se traduce en mejores respuestas; su "
        "aporte está en la calidad del contexto y de las fuentes citadas (secciones 5.3 y 5.4) y en la detección "
        "de preguntas fuera del alcance (sección 5.7), y la decisión de usarlo depende de cuánto valor se le "
        "asigne a esas propiedades frente a los 3 s adicionales."
    )
    table(
        ["Concepto", "Costo (USD)"],
        [
            ["Generación de respuestas: 4 configuraciones × 3 repeticiones × 61 preguntas", "1,30"],
            ["Juez sobre híbrida + grading (incluye pruebas de calibración de métricas)", "2,04"],
            ["Juez: completar híbrida + grading y sin RAG", "0,31"],
            ["Juez sobre solo vectorial e híbrida", "4,46"],
            ["**Total**", "**8,11**"],
        ],
        "Costo de la evaluación en la API de Gemini.",
        [12.5, 3.0],
    )
    para(
        "El costo del juez depende sobre todo de la cantidad de fragmentos: la precisión del contexto requiere una "
        "llamada por fragmento, por lo que juzgar una configuración con 7,5 fragmentos por pregunta cuesta varias veces "
        "más que juzgar la de producción."
    )


def section_errors() -> None:
    heading("6. Análisis de errores", 1)
    para("La revisión manual de las respuestas con puntaje bajo permitió identificar las siguientes causas:")
    table(
        ["Problema", "Ejemplo", "Efecto", "Posible mejora"],
        [
            ["El grading descarta todos los fragmentos", "fut-023 (expulsión en la final), fut-051", "El contexto queda vacío y el modelo responde de memoria que el Mundial 2026 aún no se jugó. Exactitud 0 con grading; 1 y 0,75 con búsqueda vectorial.", "Si el grading no aprueba ningún fragmento, usar los mejores candidatos por umbral en lugar de dejar el contexto vacío."],
            ["Información implícita entre fragmentos", "fut-024, fut-039", "Con pocos fragmentos, la respuesta usa datos que no están en el fragmento recuperado; la fidelidad cae a 0.", "Incluir el fragmento vecino o el título de la sección al seleccionar."],
            ["Fragmentos de tablas sin título", "fut-016, fut-055", "Las filas de una tabla quedan en un fragmento sin el título del documento ni de la tabla y su embedding es poco informativo.", "Anteponer el título del documento y de la tabla al fragmento al indexar."],
            ["Falsos positivos de la marca de conocimiento general", "fut-025, fut-053", "El modelo agrega una aclaración de conocimiento general aunque la respuesta proviene del contexto.", "Ajustar la instrucción del prompt sobre cuándo usar la marca."],
            ["Respuesta de memoria fuera del alcance", "fut-058", "Tras el aviso, el modelo responde con conocimiento desactualizado.", "Ante preguntas fuera del alcance, limitarse a indicar que la información no está en las fuentes."],
        ],
        "Principales causas de error identificadas.",
        [3.2, 2.6, 5.4, 4.8],
    )
    para(
        "**Causa probable del contexto vacío (fut-023 y fut-051).** En las tres repeticiones, las configuraciones sin "
        "grading recuperaron siete fragmentos que incluyen la información necesaria (la expulsión de Enzo Fernández en "
        "fut-023 y las sanciones a Verón e Independiente en fut-051). Con grading el contexto quedó vacío las tres "
        "veces, de modo que el fallo no se debe a la recuperación y es determinista (el grading usa temperatura 0). "
        "La falla ocurre entonces en el filtro: el modelo evalúa hasta 20 candidatos en una única llamada con salida "
        "estructurada, y si marca todos como no relevantes, o devuelve una lista de calificaciones vacía, el sistema "
        "no tiene un mecanismo de respaldo y pasa un contexto vacío a la generación."
    )
    para(
        "Las causas por las que el modelo rechaza los fragmentos son hipótesis que no se pudieron confirmar, porque la "
        "respuesta del evaluador no se registra. En fut-023 es posible que el modelo juzgue inverosímil un hecho de "
        "2026, posterior a su fecha de corte, o que el fragmento relevante quede diluido entre 20 candidatos largos. "
        "En fut-051, que es una pregunta de varios saltos, la respuesta está repartida en fragmentos distintos y cada "
        "uno por separado aporta solo una parte, lo que un criterio estricto (\"ayuda directamente a responder la "
        "consulta\") puede interpretar como no relevante. Para confirmarlo habría que registrar las calificaciones "
        "que devuelve el evaluador para estas dos preguntas."
    )
    para(
        "Además, durante la preparación del cuaderno de Fórmula 1 (no incluido en esta evaluación) se observó que la "
        "conversión de algunos DOCX con tablas complejas produce fragmentos con contenido corrupto, lo que señala un "
        "punto a revisar en la ingesta de ese formato."
    )


def section_conclusions() -> None:
    heading("7. Conclusiones", 1)
    bullets([
        "**P1 — Aporte del RAG:** la recuperación es la que hace útil al sistema sobre información nueva. La exactitud "
        "pasa de 0,16 a 0,82 y ninguna pregunta empeora.",
        "**P2 — Recuperación:** la configuración de producción encuentra el fragmento correcto en el 96 % de las "
        "preguntas y lo ubica primero en dos de cada tres.",
        "**P3 — Fidelidad:** las respuestas se apoyan en las fuentes (fidelidad 0,94); los casos de baja fidelidad se "
        "deben a información implícita, no a alucinaciones.",
        "**P4 — Componentes:** el componente léxico no aporta una mejora medible en este corpus. El grading mejora "
        "claramente la recuperación y la detección de preguntas fuera del alcance, pero no la exactitud final.",
        "**P5 — Fuera del alcance:** con grading, el sistema avisa en el 100 % de los casos, aunque luego el modelo "
        "puede completar con conocimiento desactualizado.",
        "**P6 — Costo:** el grading duplica la latencia mediana (de 2,5 s a 5,1 s), un costo aceptable para un chat de "
        "estudio, a cambio de un contexto más preciso y citas más confiables.",
    ])
    para(
        "En conjunto, los resultados respaldan la configuración de producción: la exactitud es equivalente a la de "
        "las alternativas más simples, pero ofrece mejores fuentes, menor contexto y una detección más confiable de lo "
        "que no está en los documentos del usuario, que son propiedades importantes en una plataforma de aprendizaje. "
        "Corregir el modo de falla del grading (descartar todos los fragmentos) es la mejora con mejor relación "
        "costo-beneficio."
    )


def section_limitations() -> None:
    heading("8. Limitaciones", 1)
    bullets([
        "**Una sola repetición del juez.** Las métricas RAGAS se calcularon sobre una repetición; el juez no es "
        "completamente determinista aun con temperatura 0. Las diferencias pequeñas (por ejemplo, 0,80 frente a 0,83 "
        "en exactitud) no deben interpretarse como significativas. Las conclusiones principales se apoyan en "
        "diferencias grandes y en la comparación pregunta por pregunta.",
        "**Un solo dominio y un tamaño acotado.** 55 preguntas dentro del alcance sobre un cuaderno de cinco "
        "documentos; el grupo de varios saltos tiene solo 10 preguntas.",
        "**Juez de la misma familia que el generador.** Juez y modelo de chat son modelos Gemini, lo que puede "
        "introducir un sesgo de autopreferencia.",
        "**Prompts del juez en inglés sobre contenido en español.** RAGAS usa prompts en inglés; se recomienda validar "
        "una muestra de puntajes con anotación humana.",
        "**Preguntas redactadas con asistencia de un LLM.** Aunque fueron revisadas manualmente, pueden reflejar el "
        "estilo del texto fuente y favorecer la búsqueda semántica.",
        "**Contradicciones en las fuentes.** Wikipedia contiene inconsistencias internas; se evitaron en las preguntas, "
        "pero el sistema puede encontrarlas en uso real.",
    ])


def section_future() -> None:
    heading("9. Trabajo futuro", 1)
    bullets([
        "Implementar un respaldo cuando el grading no aprueba ningún fragmento y volver a medir.",
        "Repetir la evaluación sobre los cuadernos de política y Fórmula 1 para comprobar si los resultados se "
        "generalizan a otros dominios y formatos.",
        "Ejecutar las repeticiones 2 y 3 del juez para reportar media y desvío estándar.",
        "Validar una muestra de puntajes con anotación humana y, si es posible, con un juez de otra familia de modelos.",
        "Evaluar variantes de top_k (4 y 12), ya definidas en el script de configuraciones.",
        "Mejorar la indexación de tablas anteponiendo el título del documento y de la tabla a cada fragmento.",
    ])


def section_reproducibility() -> None:
    heading("10. Reproducibilidad", 1)
    para("Todos los pasos se ejecutan dentro del contenedor de la API con Docker Compose:")
    table(
        ["Paso", "Comando"],
        [
            ["Descargar el corpus", "`python -m backend.eval.corpus.fetch_wikipedia`"],
            ["Exportar fragmentos indexados", "`python -m backend.eval.export_chunks --notebook-id 6ad426965b4a8f85b306b657`"],
            ["Generar respuestas", "`python -m backend.eval.run_pipeline --dataset backend/eval/datasets/futbol.jsonl --notebook-id 6ad426965b4a8f85b306b657 --configs no_rag vector_only hybrid hybrid_grading --repeats 3`"],
            ["Calcular métricas RAGAS", "`uv run --script backend/eval/evaluate_ragas.py --results-dir <carpeta> --runs 1 --configs <configuración>`"],
            ["Generar figuras", "`uv run --script backend/eval/plots.py --results-dir <carpeta>`"],
            ["Generar este informe", "`uv run --script backend/eval/informe_futbol.py`"],
        ],
        "Comandos de la evaluación.",
        [4.5, 11.5],
    )
    para(
        "Los resultados de esta corrida están en `backend/eval/results/futbol_20261018T040201Z/`: las respuestas por "
        "configuración y repetición (`*_runN.jsonl`), los puntajes del juez (`ragas/scores.jsonl`), el resumen "
        "(`ragas/summary.md`) y las figuras (`figuras/`)."
    )


def section_appendix() -> None:
    heading("Anexo. Visualizaciones alternativas", 1)
    para("Se incluyen versiones alternativas de algunas figuras para elegir la más adecuada en el capítulo de la tesis.")
    figure("02_exactitud_por_pregunta.png", "Exactitud por pregunta sin RAG frente a la configuración de producción. El tamaño de cada círculo indica la cantidad de preguntas.", 11)
    figure("02c_exactitud_slope.png", "Exactitud por pregunta al pasar de sin RAG a la configuración de producción. El grosor de cada línea indica la cantidad de preguntas; la línea punteada une las medias.", 10)
    figure("06_comparacion_por_pregunta.png", "Comparación pregunta por pregunta entre pasos consecutivos de la ablación (barras apiladas).")
    figure("07b_latencia_acumulada.png", "Distribución acumulada de la latencia por consulta y configuración (3 corridas del pipeline, todas las preguntas). El punto marca la mediana.")


def main() -> None:
    setup_styles()
    title_page()
    section_summary()
    section_objectives()
    section_system()
    section_corpus()
    section_dataset()
    section_configs()
    section_procedure()
    section_ragas()
    section_metrics()
    section_results_overview()
    section_results_rag()
    section_results_retrieval()
    section_results_ablation()
    section_results_faithfulness()
    section_results_hops()
    section_results_scope()
    section_results_cost()
    section_errors()
    section_conclusions()
    section_limitations()
    section_future()
    section_reproducibility()
    section_appendix()
    doc.save(OUTPUT_PATH)
    print(f"Informe en {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
